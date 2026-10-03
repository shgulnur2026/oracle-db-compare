#!/usr/bin/env python3
"""Сравнение таблиц территориальных БД Oracle с республиканской БД (см. prd.md).

Запуск: python app.py  ->  http://127.0.0.1:8765
"""
import csv
import datetime
import io
import json
import os
import re
import threading
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import oracledb
import openpyxl

HOST, PORT = "127.0.0.1", 8765
CHUNK = 500

REASON_MISSING = "не найдено в республиканской БД"
REASON_EMPTY = "поле {} пустое в республиканской БД"

IDENT_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_$#]{0,127}$")
FILE_RE = re.compile(r"^[A-Za-z0-9_.\-]+$")

REGION_SYNONYMS = {"район/область", "район", "область", "регион", "region"}
COLUMNS = {
    "region": REGION_SYNONYMS,
    "host": {"хост", "host"},
    "port": {"порт", "port"},
    "service": {"служба", "service", "service name", "service_name"},
    "user": {"пользователь", "user", "username"},
    "schema": {"схема", "schema"},
    "password": {"пароль", "password"},
}

TRANSLIT = dict(zip(
    "абвгдеёжзийклмнопрстуфхцчшщъыьэюяәіңғүұқөһ",
    ["a", "b", "v", "g", "d", "e", "e", "zh", "z", "i", "y", "k", "l", "m", "n", "o", "p", "r", "s",
     "t", "u", "f", "kh", "ts", "ch", "sh", "shch", "", "y", "", "e", "yu", "ya", "a", "i", "n", "g",
     "u", "u", "k", "o", "h"]))


# ----------------------------------------------------------------- состояние задания
class Job:
    def __init__(self):
        self.lock = threading.Lock()
        self.running = False
        self.stop = threading.Event()
        self.log = []
        self.regions = {}
        self.order = []
        self.summary_path = ""

    def reset(self):
        with self.lock:
            self.running = True
            self.stop = threading.Event()
            self.log = []
            self.regions = {}
            self.order = []
            self.summary_path = ""

    def add_log(self, line):
        with self.lock:
            self.log.append(line)

    def set_region(self, region, **kw):
        with self.lock:
            if region not in self.regions:
                self.regions[region] = {"region": region, "status": "ожидание", "stats": {}}
                self.order.append(region)
            r = self.regions[region]
            if "status" in kw:
                r["status"] = kw["status"]
            if "stats" in kw:
                r["stats"] = kw["stats"]
            for k in ("result_file", "log_file", "dump_file"):
                if k in kw:
                    r[k] = kw[k]

    def snapshot(self, offset):
        with self.lock:
            return {
                "running": self.running,
                "log": self.log[offset:],
                "next": len(self.log),
                "regions": [self.regions[r] for r in self.order],
                "summary": self.summary_path,
            }


JOB = Job()


class RegionLog:
    """Лог одной территории: файл + общий лог на странице. Пароли сюда не передаются."""

    def __init__(self, region, path):
        self.region = region
        self.f = open(path, "w", encoding="utf-8")

    def __call__(self, msg):
        ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.f.write(f"{ts} {msg}\n")
        self.f.flush()
        JOB.add_log(f"[{self.region}] {msg}")

    def close(self):
        self.f.close()


# ----------------------------------------------------------------- утилиты
def translit(s):
    out = "".join(TRANSLIT.get(ch, TRANSLIT.get(ch.lower(), ch)) for ch in s)
    return re.sub(r"[^A-Za-z0-9_]+", "_", out).strip("_") or "region"


def safe_name(s):
    return re.sub(r'[\\/:*?"<>|\s]+', "_", s).strip("_") or "region"


def with_suffix(filename, suffix):
    stem, ext = os.path.splitext(filename)
    return f"{stem}{suffix}{ext}"


def norm_key(v):
    if isinstance(v, Decimal):
        v = int(v) if v == v.to_integral_value() else float(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return v


def parse_list(text):
    return [t.strip() for t in re.split(r"[,;\n\r]+", text or "") if t.strip()]


def connect(c):
    return oracledb.connect(user=c["user"], password=c["password"],
                            dsn=f"{c['host']}:{c['port']}/{c['service']}")


def in_clause(keys, nk, start=1):
    """(K1,K2) IN ((:1,:2),(:3,:4),...) и плоский список параметров."""
    marks = []
    params = []
    n = start
    for k in keys:
        marks.append("(" + ",".join(f":{n + i}" for i in range(nk)) + ")")
        n += nk
        params.extend(k)
    return ",".join(marks), params


# ----------------------------------------------------------------- Excel / CSV
def read_connections_file(path):
    """Возвращает (список словарей, список ошибок)."""
    errors = []
    if not os.path.isfile(path):
        return [], [f"Файл не найден: {path}"]
    try:
        if path.lower().endswith(".csv"):
            with open(path, newline="", encoding="utf-8-sig") as f:
                sample = f.read(4096)
                f.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
                except csv.Error:
                    dialect = csv.excel
                rows = list(csv.reader(f, dialect))
        else:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
            wb.close()
    except Exception as e:
        return [], [f"Не удалось прочитать файл: {e}"]
    if not rows:
        return [], ["Файл пуст"]
    header = {}
    for i, h in enumerate(rows[0]):
        h = str(h or "").strip().lower()
        for key, syn in COLUMNS.items():
            if h in syn:
                header[key] = i
    for req, title in (("region", "Район/область"), ("host", "Хост"), ("service", "Служба")):
        if req not in header:
            errors.append(f"Нет обязательного столбца «{title}»")
    if errors:
        return [], errors
    out, seen = [], {}
    for n, row in enumerate(rows[1:], start=2):
        def cell(k):
            i = header.get(k)
            v = row[i] if i is not None and i < len(row) else None
            return "" if v is None else str(v).strip()
        if not any(str(x or "").strip() for x in row):
            continue
        rec = {k: cell(k) for k in COLUMNS}
        if rec["port"].endswith(".0"):
            rec["port"] = rec["port"][:-2]
        for req, title in (("region", "Район/область"), ("host", "Хост"), ("service", "Служба")):
            if not rec[req]:
                errors.append(f"Строка {n}: не заполнено «{title}»")
        if rec["region"]:
            if rec["region"] in seen:
                errors.append(f"Строка {n}: район «{rec['region']}» повторяет строку {seen[rec['region']]}")
            seen[rec["region"]] = n
        out.append(rec)
    return out, errors


def build_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Подключения"
    ws.append(["Район/область", "Хост", "Порт", "Служба", "Пользователь", "Схема", "Пароль"])
    ws.append(["Алматы", "10.0.0.1", 1521, "ORCL", "", "", ""])
    for col, w in zip("ABCDEFG", (22, 20, 8, 20, 18, 18, 18)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------- этап 1
def stage1_region(region, terr, repub, opts, log, stop):
    st = {"records": 0, "not_found": 0, "blob_empty": 0, "ok": 0}
    errors = 0
    result_path = os.path.join(opts["out_dir"], with_suffix(opts["result_name"], "_" + safe_name(region)))
    pk = opts["pk"]
    check = opts["check"]
    JOB.set_region(region, result_file=result_path, stats=st)
    t_conn = r_conn = None
    try:
        with open(result_path, "w", encoding="utf-8") as res:
            res.write("# REGION|TABLE|" + "|".join(pk) + "|REASON\n")
            try:
                log(f"Подключение к территориальной БД {terr['host']}:{terr['port']}/{terr['service']}")
                t_conn = connect(terr)
                log("Подключение к республиканской БД")
                r_conn = connect(repub)
            except Exception as e:
                log(f"ОШИБКА подключения: {e}")
                JOB.set_region(region, status="ошибка")
                return
            for table in opts["tables"]:
                if stop.is_set():
                    break
                try:
                    stage1_table(region, table, t_conn, r_conn, terr, repub, opts, res, st, log, stop)
                except Exception as e:
                    errors += 1
                    log(f"ОШИБКА таблицы {table}: {e}")
    finally:
        for c in (t_conn, r_conn):
            if c:
                try:
                    c.close()
                except Exception:
                    pass
    log(f"Итого: записей {st['records']}, не найдено {st['not_found']}, "
        f"{check} пуст {st['blob_empty']}, без замечаний {st['ok']}")
    if stop.is_set():
        status = "остановлено"
    elif errors:
        status = f"завершено с ошибками ({errors})"
    else:
        status = "завершено"
    JOB.set_region(region, status=status, stats=dict(st))


def stage1_table(region, table, t_conn, r_conn, terr, repub, opts, res, st, log, stop):
    pk, check = opts["pk"], opts["check"]
    nk = len(pk)
    log(f"Таблица {table}: начало")
    rc = r_conn.cursor()
    rc.execute("SELECT data_type FROM all_tab_columns WHERE owner=:o AND table_name=:t AND column_name=:c",
               o=repub["schema"].upper(), t=table.upper(), c=check.upper())
    row = rc.fetchone()
    if not row:
        raise RuntimeError(f"поле {check} или таблица не найдены в республиканской БД")
    if row[0] in ("BLOB", "CLOB", "NCLOB"):
        empty_expr = f"CASE WHEN DBMS_LOB.GETLENGTH({check})=0 THEN 1 ELSE 0 END"
    else:
        empty_expr = f"CASE WHEN LENGTH({check})=0 THEN 1 ELSE 0 END"
    # NULL (в т.ч. пустой LOB-локатор не NULL, но длина 0) -> пусто
    sel = f"SELECT {','.join(pk)}, CASE WHEN {check} IS NULL THEN 1 ELSE {empty_expr} END FROM " \
          f"{repub['schema']}.{table} WHERE ({','.join(pk)}) IN ("
    tc = t_conn.cursor()
    tc.arraysize = CHUNK
    tc.execute(f"SELECT {','.join(pk)} FROM {terr['schema']}.{table}")
    cnt = {"records": 0, "not_found": 0, "blob_empty": 0, "ok": 0}
    while not stop.is_set():
        rows = tc.fetchmany(CHUNK)
        if not rows:
            break
        keys = [tuple(norm_key(v) for v in r) for r in rows]
        marks, params = in_clause([k for k in keys if None not in k], nk)
        found = {}
        if params:
            rc.execute(sel + marks + ")", params)
            for r in rc:
                found[tuple(norm_key(v) for v in r[:nk])] = r[nk]
        lines = []
        for k in keys:
            cnt["records"] += 1
            if k not in found:
                reason, cnt["not_found"] = REASON_MISSING, cnt["not_found"] + 1
            elif found[k]:
                reason, cnt["blob_empty"] = REASON_EMPTY.format(check), cnt["blob_empty"] + 1
            else:
                cnt["ok"] += 1
                continue
            lines.append(f"{region}|{table}|" + "|".join(str(v) for v in k) + f"|{reason}\n")
            log(f"{table} {'/'.join(str(v) for v in k)}: {reason}")
        res.writelines(lines)
        res.flush()
        _add(st, cnt, region)
        cnt = {n: 0 for n in cnt}
    tc.close()
    rc.close()
    log(f"Таблица {table}: завершена")


def _add(st, cnt, region):
    for k, v in cnt.items():
        st[k] += v
    JOB.set_region(region, stats=dict(st))


# ----------------------------------------------------------------- этап 2
def read_result_keys(path, tables, nk):
    keys = {t.upper(): [] for t in tables}
    seen = {t: set() for t in keys}
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.rstrip("\r\n")
            if not line or line.startswith("#"):
                continue
            parts = line.split("|")
            if len(parts) < nk + 3:
                continue
            tbl = parts[1].upper()
            if tbl not in keys:
                continue
            k = tuple(parts[2:2 + nk])
            if k not in seen[tbl]:
                seen[tbl].add(k)
                keys[tbl].append(k)
    return keys


def coerce(value, data_type):
    if value in ("", "None"):
        return None
    if data_type in ("NUMBER", "FLOAT", "BINARY_FLOAT", "BINARY_DOUBLE", "INTEGER"):
        d = Decimal(value)
        return int(d) if d == d.to_integral_value() else float(d)
    if data_type == "DATE" or data_type.startswith("TIMESTAMP"):
        return datetime.datetime.fromisoformat(value)
    return value


def stage2_region(region, terr, opts, log, stop):
    st = {"keys": 0, "copied": 0, "not_found": 0, "anl_tables": 0}
    errors = 0
    pk = opts["pk"]
    nk = len(pk)
    result_path = os.path.join(opts["out_dir"], with_suffix(opts["result_name"], "_" + safe_name(region)))
    JOB.set_region(region, result_file=result_path, stats=st)
    if not os.path.isfile(result_path):
        log(f"ОШИБКА: файл результата этапа 1 не найден: {result_path}")
        JOB.set_region(region, status="ошибка")
        return
    keys = read_result_keys(result_path, opts["tables"], nk)
    conn = None
    created = []
    try:
        try:
            log(f"Подключение к территориальной БД {terr['host']}:{terr['port']}/{terr['service']}")
            conn = connect(terr)
        except Exception as e:
            log(f"ОШИБКА подключения: {e}")
            JOB.set_region(region, status="ошибка")
            return
        for table in opts["tables"]:
            if stop.is_set():
                break
            try:
                n = stage2_table(region, table, keys[table.upper()], conn, terr, opts, st, log, stop)
                if n:
                    created.append(table + "ANL")
                    st["anl_tables"] += 1
            except Exception as e:
                errors += 1
                log(f"ОШИБКА таблицы {table}: {e}")
            JOB.set_region(region, stats=dict(st))
        if created and not stop.is_set():
            dump = with_suffix(opts["dump_name"], "_" + translit(region))
            try:
                run_datapump(conn, terr["schema"], created, opts["directory"], dump, log)
                JOB.set_region(region, dump_file=f"{opts['directory']}:{dump}")
            except Exception as e:
                errors += 1
                log(f"ОШИБКА экспорта Data Pump: {e}")
        elif not created and not stop.is_set():
            log("Нет созданных таблиц ANL — дамп не создаётся")
    finally:
        if conn:
            try:
                conn.close()
            except Exception:
                pass
    log(f"Итого: ключей {st['keys']}, скопировано {st['copied']}, не найдено в БД {st['not_found']}, "
        f"таблиц ANL {st['anl_tables']}")
    if stop.is_set():
        status = "остановлено"
    elif errors:
        status = f"завершено с ошибками ({errors})"
    else:
        status = "завершено"
    JOB.set_region(region, status=status, stats=dict(st))


def stage2_table(region, table, raw_keys, conn, terr, opts, st, log, stop):
    """Создаёт таблицу ANL и копирует записи. Возвращает True, если таблица создана."""
    pk = opts["pk"]
    nk = len(pk)
    schema = terr["schema"]
    anl = table + "ANL"
    if len(anl) > 128:
        raise RuntimeError(f"имя {anl} слишком длинное")
    st["keys"] += len(raw_keys)
    if not raw_keys:
        log(f"Таблица {table}: в файле результата нет ключей, пропущена")
        return False
    cur = conn.cursor()
    cur.execute("SELECT column_name, data_type FROM all_tab_columns WHERE owner=:o AND table_name=:t",
                o=schema.upper(), t=table.upper())
    types = dict(cur.fetchall())
    if not types:
        raise RuntimeError("таблица не найдена")
    for p in pk:
        if p.upper() not in types:
            raise RuntimeError(f"поле ключа {p} не найдено")
    keys = [tuple(coerce(v, types[p.upper()]) for v, p in zip(k, pk)) for k in raw_keys]
    cur.execute("SELECT COUNT(*) FROM all_tables WHERE owner=:o AND table_name=:t",
                o=schema.upper(), t=anl.upper())
    if cur.fetchone()[0]:
        if not opts["recreate"]:
            raise RuntimeError(f"таблица {anl} уже существует (включите пересоздание)")
        log(f"Удаление существующей таблицы {anl}")
        cur.execute(f"DROP TABLE {schema}.{anl} PURGE")
    log(f"Создание таблицы {anl}, ключей {len(keys)}")
    cur.execute(f"CREATE TABLE {schema}.{anl} AS SELECT * FROM {schema}.{table} WHERE 1=0")
    copied = 0
    for i in range(0, len(keys), CHUNK):
        if stop.is_set():
            break
        marks, params = in_clause(keys[i:i + CHUNK], nk)
        cur.execute(f"INSERT INTO {schema}.{anl} SELECT * FROM {schema}.{table} "
                    f"WHERE ({','.join(pk)}) IN ({marks})", params)
        copied += cur.rowcount
    conn.commit()
    st["copied"] += copied
    st["not_found"] += max(0, len(keys) - copied)
    log(f"Таблица {anl}: скопировано {copied} из {len(keys)}")
    cur.close()
    return True


def run_datapump(conn, schema, tables, directory, dump, log):
    log(f"Экспорт Data Pump: {len(tables)} таблиц(ы) -> {directory}:{dump}")
    names = ",".join(f"''{t.upper()}''" for t in tables)
    plsql = f"""
DECLARE
  h NUMBER;
  js VARCHAR2(30);
BEGIN
  h := DBMS_DATAPUMP.OPEN('EXPORT', 'TABLE', NULL, :job);
  DBMS_DATAPUMP.ADD_FILE(h, :dump, :dir, NULL, DBMS_DATAPUMP.KU$_FILE_TYPE_DUMP_FILE, 1);
  DBMS_DATAPUMP.ADD_FILE(h, :lg, :dir, NULL, DBMS_DATAPUMP.KU$_FILE_TYPE_LOG_FILE, 1);
  DBMS_DATAPUMP.METADATA_FILTER(h, 'SCHEMA_EXPR', 'IN (''{schema.upper()}'')');
  DBMS_DATAPUMP.METADATA_FILTER(h, 'NAME_EXPR', 'IN ({names})');
  DBMS_DATAPUMP.START_JOB(h);
  DBMS_DATAPUMP.WAIT_FOR_JOB(h, js);
  :state := js;
END;"""
    job = "CMP_ANL_" + datetime.datetime.now().strftime("%H%M%S%f")[:10]
    state = conn.cursor().var(str)
    conn.cursor().execute(plsql, job=job, dump=dump, dir=directory.upper(),
                          lg=os.path.splitext(dump)[0] + ".log", state=state)
    if state.getvalue() != "COMPLETED":
        raise RuntimeError(f"задание Data Pump завершилось со статусом {state.getvalue()}")
    log(f"Дамп создан на сервере БД: {directory}:{dump}")


# ----------------------------------------------------------------- запуск
def validate(p):
    """Проверка параметров. Возвращает (opts, список ошибок)."""
    errs = []

    def need_ident(value, title):
        if not IDENT_RE.match(value or ""):
            errs.append(f"{title}: недопустимое или пустое имя «{value}»")

    stage = int(p.get("stage", 1))
    multi = p.get("mode") == "multi"
    opts = {
        "stage": stage, "multi": multi,
        "tables": parse_list(p.get("tables")),
        "pk": [x.strip() for x in parse_list(p.get("pk")) or ["ID", "KNR"]],
        "check": (p.get("check") or "BlobArch").strip(),
        "out_dir": (p.get("out_dir") or "").strip(),
        "result_name": (p.get("result_name") or "compare_result.txt").strip(),
        "log_name": (p.get("log_name") or "compare_log.log").strip(),
        "summary_name": (p.get("summary_name") or "compare_summary.txt").strip(),
        "directory": (p.get("directory") or "").strip(),
        "dump_name": (p.get("dump_name") or "compare_anl.dmp").strip(),
        "recreate": bool(p.get("recreate")),
        "threads": max(1, min(32, int(p.get("threads") or 4))),
    }
    if not opts["tables"]:
        errs.append("Не указан список таблиц")
    for t in opts["tables"]:
        need_ident(t, "Таблица")
    for k in opts["pk"]:
        need_ident(k, "Поле ключа")
    if stage == 1:
        need_ident(opts["check"], "Поле для проверки")
    if not opts["out_dir"]:
        errs.append("Не указана папка результатов")
    for key, title in (("result_name", "Имя файла результата"), ("log_name", "Имя файла лога"),
                       ("summary_name", "Имя файла сводки")):
        if not opts[key] or re.search(r'[\\/:*?"<>|]', opts[key]):
            errs.append(f"{title}: недопустимое имя")
    if stage == 2:
        need_ident(opts["directory"], "Каталог Oracle (DIRECTORY)")
        if not FILE_RE.match(opts["dump_name"]):
            errs.append("Имя файла дампа: допустимы только латиница, цифры, _ . -")

    def conn_of(prefix, title, need_target=True):
        c = {k: (p.get(f"{prefix}_{k}") or "").strip() for k in ("host", "port", "service", "user", "schema")}
        c["password"] = p.get(f"{prefix}_password") or ""
        c["port"] = c["port"] or "1521"
        if need_target:
            for k, n in (("host", "хост"), ("service", "служба"), ("user", "пользователь"),
                         ("schema", "схема")):
                if not c[k]:
                    errs.append(f"{title}: не заполнено поле «{n}»")
        if c["schema"]:
            need_ident(c["schema"], f"{title}: схема")
        return c

    repub = None
    if stage == 1:
        repub = conn_of("repub", "Республиканская БД")
    defaults = conn_of("terr", "Территориальная БД", need_target=not multi)
    terrs = []
    if multi:
        path = (p.get("excel_path") or "").strip()
        if not path:
            errs.append("Не указан файл со списком БД")
        else:
            rows, e = read_connections_file(path)
            errs.extend(e)
            for r in rows:
                terrs.append((r["region"], {
                    "host": r["host"], "service": r["service"],
                    "port": r["port"] or defaults["port"],
                    "user": r["user"] or defaults["user"],
                    "schema": r["schema"] or defaults["schema"],
                    "password": r["password"] or defaults["password"],
                }))
            for region, c in terrs:
                if not c["user"] or not c["schema"]:
                    errs.append(f"{region}: не заполнены пользователь/схема (ни в файле, ни на странице)")
                elif not IDENT_RE.match(c["schema"]):
                    errs.append(f"{region}: недопустимое имя схемы")
            if not terrs and not e:
                errs.append("В файле нет ни одной БД")
    else:
        region = (p.get("region") or "").strip()
        if not region:
            errs.append("Не указано поле «Район/область»")
        terrs = [(region, defaults)]
    opts["repub"], opts["terrs"] = repub, terrs
    return opts, errs


def run_job(opts):
    stage, stop = opts["stage"], JOB.stop
    os.makedirs(opts["out_dir"], exist_ok=True)

    def work(item):
        region, terr = item
        JOB.set_region(region, status="ожидание")
        if stop.is_set():
            JOB.set_region(region, status="пропущено (остановлено)")
            return
        suffix = ("_anl_" if stage == 2 else "_") + safe_name(region)
        log_path = os.path.join(opts["out_dir"], with_suffix(opts["log_name"], suffix))
        log = RegionLog(region, log_path)
        JOB.set_region(region, status="выполняется", log_file=log_path)
        try:
            if stage == 1:
                stage1_region(region, terr, opts["repub"], opts, log, stop)
            else:
                stage2_region(region, terr, opts, log, stop)
        except Exception as e:
            log(f"ОШИБКА: {e}")
            JOB.set_region(region, status="ошибка")
        finally:
            log.close()

    try:
        for region, _ in opts["terrs"]:
            JOB.set_region(region, status="ожидание")
        with ThreadPoolExecutor(max_workers=opts["threads"] if opts["multi"] else 1) as ex:
            list(ex.map(work, opts["terrs"]))
        if opts["multi"]:
            write_summary(opts)
        JOB.add_log("Обработка завершена" if not stop.is_set() else "Обработка остановлена")
    except Exception as e:
        JOB.add_log(f"ОШИБКА: {e}")
    finally:
        with JOB.lock:
            JOB.running = False


def write_summary(opts):
    name = opts["summary_name"]
    if opts["stage"] == 2:
        name = with_suffix(name, "_anl")
    path = os.path.join(opts["out_dir"], name)
    if opts["stage"] == 1:
        head = "REGION|STATUS|RECORDS|NOT_FOUND|BLOB_EMPTY|OK|RESULT_FILE|LOG_FILE"
        keys = ("records", "not_found", "blob_empty", "ok")
    else:
        head = "REGION|STATUS|KEYS|COPIED|NOT_FOUND|ANL_TABLES|DUMP|LOG_FILE"
        keys = ("keys", "copied", "not_found", "anl_tables")
    with open(path, "w", encoding="utf-8") as f:
        f.write("# " + head + "\n")
        for r in JOB.snapshot(0)["regions"]:
            nums = [str(r["stats"].get(k, 0)) for k in keys]
            third = r.get("dump_file", "") if opts["stage"] == 2 else r.get("result_file", "")
            f.write("|".join([r["region"], r["status"], *nums, third, r.get("log_file", "")]) + "\n")
    with JOB.lock:
        JOB.summary_path = path
    JOB.add_log(f"Сводка сохранена: {path}")


# ----------------------------------------------------------------- веб-сервер
class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):  # не пишем запросы (там могут быть пароли)
        pass

    def _send(self, code, body, ctype="application/json; charset=utf-8", headers=None):
        if isinstance(body, (dict, list)):
            body = json.dumps(body, ensure_ascii=False)
        if isinstance(body, str):
            body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        for k, v in (headers or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path, _, q = self.path.partition("?")
        if path == "/":
            self._send(200, INDEX_HTML, "text/html; charset=utf-8")
        elif path == "/status":
            m = re.search(r"offset=(\d+)", q)
            self._send(200, JOB.snapshot(int(m.group(1)) if m else 0))
        elif path == "/template.xlsx":
            self._send(200, build_template(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       {"Content-Disposition": 'attachment; filename="databases_template.xlsx"'})
        else:
            self._send(404, {"error": "not found"})

    def do_POST(self):
        n = int(self.headers.get("Content-Length") or 0)
        try:
            p = json.loads(self.rfile.read(n) or b"{}")
        except ValueError:
            return self._send(400, {"errors": ["Некорректный запрос"]})
        if self.path == "/stop":
            JOB.stop.set()
            JOB.add_log("Запрошена остановка (после текущей порции данных)")
            return self._send(200, {"ok": True})
        if self.path == "/start":
            if JOB.running:
                return self._send(409, {"errors": ["Процесс уже запущен"]})
            try:
                opts, errs = validate(p)
            except Exception as e:
                return self._send(400, {"errors": [f"Ошибка параметров: {e}"]})
            if errs:
                return self._send(400, {"errors": errs})
            JOB.reset()
            threading.Thread(target=run_job, args=(opts,), daemon=True).start()
            return self._send(200, {"ok": True})
        self._send(404, {"error": "not found"})


INDEX_HTML = r"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><title>Сравнение БД Oracle</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
body{font:14px system-ui,sans-serif;margin:0;background:#f4f5f7;color:#222}
main{max-width:1100px;margin:0 auto;padding:16px}
h1{font-size:20px} fieldset{background:#fff;border:1px solid #d5d8dd;border-radius:6px;margin:0 0 12px;padding:10px 14px}
legend{font-weight:600;padding:0 6px}
.g{display:grid;grid-template-columns:repeat(auto-fill,minmax(210px,1fr));gap:8px}
label{display:flex;flex-direction:column;font-size:12px;color:#555;gap:2px}
input,textarea,select{font:inherit;padding:5px;border:1px solid #bbb;border-radius:4px}
.row{display:flex;gap:18px;align-items:center;flex-wrap:wrap;margin-bottom:6px}
.row label{flex-direction:row;align-items:center;font-size:14px;color:#222;gap:5px}
button{font:inherit;padding:7px 18px;border:0;border-radius:4px;cursor:pointer;color:#fff;background:#2563eb}
button.stop{background:#c0392b} button:disabled{opacity:.5;cursor:default}
#log{background:#111;color:#cfe;font:12px/1.4 Consolas,monospace;height:260px;overflow:auto;white-space:pre-wrap;padding:8px;border-radius:4px}
table{border-collapse:collapse;width:100%;background:#fff;font-size:13px}
th,td{border:1px solid #d5d8dd;padding:4px 8px;text-align:left} th{background:#eef0f3}
#err{color:#b00020;white-space:pre-wrap;margin:6px 0}
.hide{display:none}
</style></head><body><main>
<h1>Сравнение таблиц территориальных БД с республиканской БД</h1>
<fieldset><legend>Режим</legend><div class="row">
<label>Этап <select id="stage"><option value="1">1 — сравнение</option><option value="2">2 — таблицы ANL и дамп</option></select></label>
<label>Подключение <select id="mode"><option value="single">одна территориальная БД</option><option value="multi">множество БД (Excel/CSV)</option></select></label>
</div></fieldset>

<fieldset id="fs-repub"><legend>Республиканская БД</legend><div class="g">
<label>Хост<input id="repub_host"></label><label>Порт<input id="repub_port" value="1521"></label>
<label>Служба (service name)<input id="repub_service"></label><label>Пользователь<input id="repub_user"></label>
<label>Схема<input id="repub_schema"></label><label>Пароль<input id="repub_password" type="password" autocomplete="off"></label>
</div></fieldset>

<fieldset><legend id="terr-title">Территориальная БД</legend>
<div class="g">
<label id="l-region">Район/область<input id="region"></label>
<label class="s">Хост<input id="terr_host"></label><label>Порт<input id="terr_port" value="1521"></label>
<label class="s">Служба (service name)<input id="terr_service"></label><label>Пользователь<input id="terr_user"></label>
<label>Схема<input id="terr_schema"></label><label>Пароль<input id="terr_password" type="password" autocomplete="off"></label>
</div>
<div id="multi-box" class="hide" style="margin-top:8px"><div class="g">
<label>Путь к файлу .xlsx / .csv<input id="excel_path"></label>
<label>Потоков (1–32)<input id="threads" type="number" min="1" max="32" value="4"></label>
<label>Имя файла сводки<input id="summary_name" value="compare_summary.txt"></label></div>
<p><a href="/template.xlsx">Скачать шаблон Excel</a>. Пустые поля порт/пользователь/схема/пароль берутся из формы выше.</p></div>
</fieldset>

<fieldset><legend>Что обрабатывать</legend><div class="g">
<label style="grid-column:span 2">Таблицы (через запятую или с новой строки)<textarea id="tables" rows="3"></textarea></label>
<label>Поля первичного ключа<input id="pk" value="ID, KNR"></label>
<label id="l-check">Поле для проверки<input id="check" value="BlobArch"></label>
</div></fieldset>

<fieldset><legend>Файлы</legend><div class="g">
<label>Папка результатов<input id="out_dir"></label>
<label>Файл результата<input id="result_name" value="compare_result.txt"></label>
<label>Файл лога<input id="log_name" value="compare_log.log"></label>
</div></fieldset>

<fieldset id="fs-s2" class="hide"><legend>Этап 2</legend><div class="g">
<label>Каталог Oracle (DIRECTORY)<input id="directory"></label>
<label>Имя файла дампа<input id="dump_name" value="compare_anl.dmp"></label>
<label class="row" style="flex-direction:row;align-items:center"><input type="checkbox" id="recreate"> Пересоздавать существующие таблицы ANL</label>
</div></fieldset>

<div class="row"><button id="start">Запустить</button><button id="stop" class="stop" disabled>Остановить</button><span id="state"></span></div>
<div id="err"></div>
<h3>Результаты по районам</h3>
<table><thead id="thead"></thead><tbody id="tbody"></tbody></table>
<h3>Лог</h3><div id="log"></div>
</main>
<script>
const $=id=>document.getElementById(id);
const FIELDS=["stage","mode","repub_host","repub_port","repub_service","repub_user","repub_schema",
"region","terr_host","terr_port","terr_service","terr_user","terr_schema","excel_path","threads","summary_name",
"tables","pk","check","out_dir","result_name","log_name","directory","dump_name"];
for(const f of FIELDS){try{const v=localStorage.getItem("cmp_"+f);if(v!==null)$(f).value=v;}catch(e){}
 $(f).addEventListener("input",()=>{try{localStorage.setItem("cmp_"+f,$(f).value)}catch(e){}});}
try{$("recreate").checked=localStorage.getItem("cmp_recreate")==="1"}catch(e){}
$("recreate").onchange=()=>{try{localStorage.setItem("cmp_recreate",$("recreate").checked?"1":"0")}catch(e){}};
function layout(){
 const s2=$("stage").value==="2",multi=$("mode").value==="multi";
 $("fs-repub").classList.toggle("hide",s2);$("fs-s2").classList.toggle("hide",!s2);
 $("l-check").classList.toggle("hide",s2);$("multi-box").classList.toggle("hide",!multi);
 $("l-region").classList.toggle("hide",multi);
 document.querySelectorAll("label.s").forEach(l=>l.classList.toggle("hide",multi));
 $("terr-title").textContent=multi?"Параметры по умолчанию для территориальных БД":"Территориальная БД";
 $("thead").innerHTML=s2?"<tr><th>Район</th><th>Статус</th><th>Ключей в файле</th><th>Скопировано в ANL</th><th>Не найдено в БД</th><th>Таблиц ANL</th></tr>"
  :"<tr><th>Район</th><th>Статус</th><th>Записей</th><th>Не найдено</th><th>BlobArch пуст</th><th>Без замечаний</th></tr>";
}
$("stage").onchange=$("mode").onchange=layout;layout();
function payload(){const p={recreate:$("recreate").checked};
 for(const el of document.querySelectorAll("input,select,textarea")){if(el.id&&el.type!=="checkbox")p[el.id]=el.value;}return p;}
let offset=0,timer=null;
async function start(){
 $("err").textContent="";
 const r=await fetch("/start",{method:"POST",body:JSON.stringify(payload())});
 const j=await r.json();
 if(!r.ok){$("err").textContent=(j.errors||[]).join("\n");return;}
 offset=0;$("log").textContent="";poll();
}
async function poll(){
 clearTimeout(timer);
 const j=await (await fetch("/status?offset="+offset)).json();
 offset=j.next;
 if(j.log.length){const l=$("log");l.textContent+=j.log.join("\n")+"\n";l.scrollTop=l.scrollHeight;}
 const s2=$("stage").value==="2",keys=s2?["keys","copied","not_found","anl_tables"]:["records","not_found","blob_empty","ok"];
 $("tbody").innerHTML=j.regions.map(r=>"<tr><td>"+esc(r.region)+"</td><td>"+esc(r.status)+"</td>"+
  keys.map(k=>"<td>"+(r.stats[k]??"")+"</td>").join("")+"</tr>").join("");
 $("start").disabled=j.running;$("stop").disabled=!j.running;$("state").textContent=j.running?"выполняется…":"";
 if(j.running)timer=setTimeout(poll,1000);
}
function esc(s){return String(s).replace(/[&<>]/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;"}[c]))}
$("start").onclick=start;
$("stop").onclick=()=>fetch("/stop",{method:"POST",body:"{}"});
poll();
</script></body></html>
"""


def main():
    srv = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Откройте http://{HOST}:{PORT}")
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
